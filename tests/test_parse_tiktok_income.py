"""TikTok Shop income file (income_*.xlsx) -> parse_tiktok_income.

Plan: projects/tiktok-order-import/tiktok-order-import-plan.md, PR-2 step 1.
The fixture is the real income file of 2026-09-30 (3 orders, no buyer data;
bank column '/'). Oracle = the plan's step 5: payouts 309.05 / 88.38 / 0
(sum 397.43 == sheet รายงาน), fees 125.95 / 30.62 / 0.

Three checks refuse the whole file (fail loud, nothing written):
  (a) per row, the leaf fee buckets sum to ค่าธรรมเนียมทั้งหมด;
  (b) per row, รายได้ทั้งหมด + ค่าธรรมเนียมทั้งหมด + จำนวนการปรับยอด ==
      ยอดการชำระเงินทั้งหมด (TikTok's own identity; holds on all 3 sample rows);
  (c) the payouts of ALL rows sum to the รายงาน sheet's ยอดการชำระเงินทั้งหมด.
"""
import io
import json
import os

import openpyxl
import pandas as pd
import pytest

from parse_tiktok_income import parse_tiktok_income, load_tiktok_income, TikTokIncomeError

FIXTURE = os.path.join(os.path.dirname(__file__), 'fixtures', 'tiktok', 'tiktok_income_sample.xlsx')
DETAIL = 'รายละเอียดคำสั่งซื้อ'
REPORT = 'รายงาน'

O379, O543, O817 = '585884671861360379', '585884215723460543', '585883444661159817'


def _bytes(edit=None):
    """The fixture, optionally edited cell by cell through openpyxl (never
    read_only: it stops after 1 data row on this file)."""
    with open(FIXTURE, 'rb') as f:
        data = f.read()
    if edit is None:
        return data
    wb = openpyxl.load_workbook(io.BytesIO(data))
    edit(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _col(ws, header):
    for c in ws[1]:
        if c.value == header:
            return c.column
    raise AssertionError(f'no column {header!r}')


def _set(ws, row, header, value):
    ws.cell(row=row, column=_col(ws, header), value=value)


def _parse(data=None):
    return parse_tiktok_income(*load_tiktok_income(io.BytesIO(data or _bytes())))


def test_all_three_detail_rows_are_read():
    detail, report = load_tiktok_income(io.BytesIO(_bytes()))
    assert len(detail) == 3
    assert list(detail['หมายเลขคำสั่งซื้อ/การปรับ']) == [O379, O543, O817]


def test_settlements_are_the_payouts():
    p = _parse()
    s = {x['order_sn']: x for x in p['settlements']}
    assert [s[o]['actual_payout'] for o in (O379, O543, O817)] == [309.05, 88.38, 0.0]
    assert s[O379]['settled_at'] == '2026-09-13'
    assert s[O817]['settled_at'] == '2026-09-05'
    assert p['adjustments'] == []
    assert p['total'] == 397.43


def test_fee_rows_follow_the_shopee_identity():
    f = {x['order_sn']: x for x in _parse()['fee_rows']}
    assert [f[o]['fee_total'] for o in (O379, O543, O817)] == [125.95, 30.62, 0.0]
    assert [f[o]['item_value'] for o in (O379, O543, O817)] == [435.0, 119.0, 0.0]
    assert [f[o]['net_payout'] for o in (O379, O543, O817)] == [309.05, 88.38, 0.0]


def test_buckets_for_379():
    r = {x['order_sn']: x for x in _parse()['fee_rows']}[O379]
    assert (r['fee_commission'], r['fee_transaction'], r['fee_ads_escrow'], r['fee_service'],
            r['fee_platform'], r['shipping_net']) == (-41.89, -13.96, -36.45, -32.58, -1.07, 0.0)
    raw = json.loads(r['fee_raw_json'])
    assert raw['ค่าคอมมิชชั่น TikTok Shop'] == -41.89
    assert raw['ธนาคารของลูกค้าสำหรับการชำระเงิน'] is None
    assert raw['หมายเลขคำสั่งซื้อ/การปรับ'] == O379          # text, not a rounded float


def test_adjustment_row_is_skipped_and_named():
    def edit(wb):
        ws = wb[DETAIL]
        _set(ws, 4, 'ประเภทธุรกรรม', 'การปรับยอด')
    p = _parse(_bytes(edit))
    assert [x['order_sn'] for x in p['settlements']] == [O379, O543]
    assert [x['order_sn'] for x in p['fee_rows']] == [O379, O543]
    assert p['adjustments'] == [{'id': O817, 'type': 'การปรับยอด', 'amount': 0.0}]


# ── (a) buckets == ค่าธรรมเนียมทั้งหมด ──────────────────────────────────────────

def test_check_a_refuses_when_buckets_miss_the_total():
    def edit(wb):
        _set(wb[DETAIL], 2, 'ค่าธรรมเนียมโครงสร้างพื้นฐาน', '-2.07')
        _set(wb[DETAIL], 2, 'ค่าธรรมเนียมทั้งหมด', '-125.95')   # total not moved
    with pytest.raises(TikTokIncomeError, match=O379):
        _parse(_bytes(edit))


# ── (b) รายได้ + ค่าธรรมเนียม + การปรับยอด == ยอดการชำระเงิน ─────────────────────────

def test_check_b_refuses_when_the_row_identity_breaks():
    def edit(wb):
        _set(wb[DETAIL], 3, 'รายได้ทั้งหมด', '120')
    with pytest.raises(TikTokIncomeError, match=O543):
        _parse(_bytes(edit))


def test_check_b_counts_the_adjustment_amount():
    """Control for (b): moving the payout and the adjustment together keeps it."""
    def edit(wb):
        ws = wb[DETAIL]
        _set(ws, 3, 'จำนวนการปรับยอด', '10')
        _set(ws, 3, 'ยอดการชำระเงินทั้งหมด', '98.38')
        ws2 = wb[REPORT]
        for row in ws2.iter_rows():
            if row[1].value == 'ยอดการชำระเงินทั้งหมด':
                row[5].value = '407.43'
    p = _parse(_bytes(edit))
    assert p['total'] == 407.43


# ── (c) Σ payouts == รายงาน total ───────────────────────────────────────────────

def test_check_c_refuses_when_the_report_total_disagrees():
    def edit(wb):
        for row in wb[REPORT].iter_rows():
            if row[1].value == 'ยอดการชำระเงินทั้งหมด':
                row[5].value = '397.44'
    with pytest.raises(TikTokIncomeError, match='397.44'):
        _parse(_bytes(edit))


def test_check_c_counts_adjustment_rows_too():
    """An adjustment row is not imported, but its money is in the report total."""
    def edit(wb):
        ws = wb[DETAIL]
        _set(ws, 4, 'ประเภทธุรกรรม', 'การปรับยอด')
        _set(ws, 4, 'จำนวนการปรับยอด', '5')
        _set(ws, 4, 'ยอดการชำระเงินทั้งหมด', '5')
        for row in wb[REPORT].iter_rows():
            if row[1].value == 'ยอดการชำระเงินทั้งหมด':
                row[5].value = '402.43'
    p = _parse(_bytes(edit))
    assert p['adjustments'] == [{'id': O817, 'type': 'การปรับยอด', 'amount': 5.0}]
    assert p['total'] == 402.43


# ── unknown columns, dates, rows ─────────────────────────────────────────────────

def test_unknown_nonzero_column_is_refused_by_name():
    def edit(wb):
        ws = wb[DETAIL]
        c = ws.max_column + 1
        ws.cell(row=1, column=c, value='ค่าธรรมเนียมใหม่')
        ws.cell(row=2, column=c, value='-3')
    with pytest.raises(TikTokIncomeError, match='ค่าธรรมเนียมใหม่') as e:
        _parse(_bytes(edit))
    # The message names both homes a new column can have.
    assert 'TIKTOK_BUCKET' in str(e.value) and '_OTHER_COLUMNS' in str(e.value)


def test_unknown_all_zero_column_is_accepted():
    def edit(wb):
        ws = wb[DETAIL]
        c = ws.max_column + 1
        ws.cell(row=1, column=c, value='ค่าธรรมเนียมใหม่')
        for r in (2, 3, 4):
            ws.cell(row=r, column=c, value='0')
    assert len(_parse(_bytes(edit))['fee_rows']) == 3


def test_unparseable_settled_date_is_refused():
    def edit(wb):
        _set(wb[DETAIL], 2, 'เวลาที่ชำระคำสั่งซื้อ', '13-09-2026')
    with pytest.raises(TikTokIncomeError, match='13-09-2026'):
        _parse(_bytes(edit))


def test_missing_sheet_is_refused():
    def edit(wb):
        del wb[REPORT]
    with pytest.raises(TikTokIncomeError):
        _parse(_bytes(edit))


def _shift(wb, row, col, delta, total_too=True):
    """Add `delta` to one fee cell of `row`, and to the row's totals so (a) (b) (c) hold."""
    ws = wb[DETAIL]
    def bump(header, d):
        c = ws.cell(row=row, column=_col(ws, header))
        c.value = str(round(float(c.value) + d, 2))
    bump(col, delta)
    if total_too:
        bump('ค่าธรรมเนียมทั้งหมด', delta)
        bump('ยอดการชำระเงินทั้งหมด', delta)
        for r in wb[REPORT].iter_rows():
            if r[1].value == 'ยอดการชำระเงินทั้งหมด':
                r[5].value = str(round(float(r[5].value) + delta, 2))


def test_affiliate_siblings_are_leaves_in_the_affiliate_bucket():
    """The รายงาน sheet puts ค่าคอมมิชชั่นของพาร์ทเนอร์แอฟฟิลิเอต BESIDE
    ค่าคอมมิชชั่นแอฟฟิลิเอต, not under it: a leaf, counted once."""
    def edit(wb):
        _shift(wb, 3, 'ค่าคอมมิชชั่นของพาร์ทเนอร์แอฟฟิลิเอต', -2.0)
    r = {x['order_sn']: x for x in _parse(_bytes(edit))['fee_rows']}[O543]
    assert r['fee_ads_escrow'] == -7.36
    assert r['fee_total'] == 32.62


def test_a_subtotal_part_is_never_added_again():
    """A part under ค่าคอมมิชชั่นแอฟฟิลิเอต moves its subtotal too; the bucket
    takes the subtotal only, so the part is not counted twice."""
    def edit(wb):
        _shift(wb, 3, 'ภาษีเงินได้บุคคลธรรมดาที่หักไว้จากคอมมิชชั่นของแอฟฟิลิเอต', -1.0,
               total_too=False)
        _shift(wb, 3, 'ค่าคอมมิชชั่นแอฟฟิลิเอต', -1.0)
    r = {x['order_sn']: x for x in _parse(_bytes(edit))['fee_rows']}[O543]
    assert r['fee_ads_escrow'] == -6.36


def test_check_b_skips_adjustment_rows():
    """Adjustment rows are not imported and their money is pinned by (c); the
    sample has none, so (b) must not refuse a file over a shape never seen."""
    def edit(wb):
        ws = wb[DETAIL]
        _set(ws, 4, 'ประเภทธุรกรรม', 'การปรับยอด')
        _set(ws, 4, 'ยอดการชำระเงินทั้งหมด', '5')      # revenue/fees/adjustment stay 0
        for row in wb[REPORT].iter_rows():
            if row[1].value == 'ยอดการชำระเงินทั้งหมด':
                row[5].value = '402.43'
    p = _parse(_bytes(edit))
    assert p['adjustments'] == [{'id': O817, 'type': 'การปรับยอด', 'amount': 5.0}]
    assert p['total'] == 402.43
